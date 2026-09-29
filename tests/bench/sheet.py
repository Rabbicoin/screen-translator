# -*- coding: utf-8 -*-
"""Лист «Проверка валют»: столбцы I–J — прогон каждого случая через пересчёт цен.

Клетка рисуется в headless Chrome шрифтом Calibri 14 пт при масштабе 100 %, как
в Excel, и проходит тот же путь, что у Ctrl+Alt+D: read_lines -> find_prices.
I — язык распознавания «eng», J — «авто», как у новых пользователей (языки по
письменности снимка плюс английский).

Запуск:  python tests/bench/sheet.py "tests/Книга1(АвтоматическиВосстановлено).xlsx" [подпись]
Подпись попадает в заголовки I–J, например «29.09». Копия книги до записи —
во временной папке Windows, путь печатается.
"""
import os
import shutil
import subprocess
import sys
import tempfile
from concurrent.futures import ProcessPoolExecutor

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.environ.get("REPO", os.path.dirname(os.path.dirname(HERE)))
ARGS = sys.argv[1:]
sys.argv = sys.argv[:1]
sys.path.insert(0, HERE)

from cases import SECTIONS  # noqa: E402
from bench import CHROME, SHOTS  # noqa: E402

CASES = [c for _, section in SECTIONS for c in section]
CELL_W, CELL_H = 440, 60


def shoot():
    """Все случаи столбиком на одной странице -> путь к снимку."""
    cells = []
    for i, (text, _src, _want, style, _note) in enumerate(CASES):
        t = text.replace("&", "&amp;").replace("<", "&lt;")
        if style == "strike":                       # первая цена зачёркнута
            first, rest = t.split(" ", 1)
            t = f"<s>{first}</s> {rest}"
        cells.append(f'<div style="top:{i * CELL_H}px">{t}</div>')
    html = ("<!doctype html><meta charset=utf-8><style>body{margin:0;background:#fff}"
            f"div{{position:absolute;left:0;width:{CELL_W - 12}px;height:{CELL_H}px;"
            f"padding-left:12px;line-height:{CELL_H}px;font:14pt Calibri;color:#000;"
            "white-space:nowrap}</style>" + "".join(cells))
    os.makedirs(SHOTS, exist_ok=True)
    hp = os.path.join(SHOTS, "sheet_calibri14.html")
    png = hp[:-5] + ".png"
    open(hp, "w", encoding="utf-8").write(html)
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--hide-scrollbars",
                    "--force-device-scale-factor=1",
                    f"--window-size={CELL_W},{len(CASES) * CELL_H}",
                    f"--screenshot={png}", "file:///" + hp.replace("\\", "/")],
                   check=True, capture_output=True, timeout=120)
    return png


_ST = None


def _init():
    global _ST
    sys.path.insert(0, REPO)
    import screen_translator  # noqa: F401  (tesseract и языки)
    screen_translator.CFG["ocr_langs"] = "auto"
    _ST = screen_translator


def _amount(a):
    return ("%f" % a).rstrip("0").rstrip(".")


def work(job):
    png, i, mode = job
    from PIL import Image
    import currency
    text, source, want, _style, _note = CASES[i]
    crop = Image.open(png).convert("RGB").crop((0, i * CELL_H, CELL_W, (i + 1) * CELL_H))
    langs = "eng" if mode == "eng" else _ST.price_ocr_langs(crop)
    lines = currency.read_lines(crop, langs, 2.0)
    got = [(round(p.amount, 6), p.code) for p in currency.find_prices(lines, source)]
    if got == [(round(a, 6), c) for a, c in want]:
        return "ок"
    read = " | ".join(" ".join(w["text"] for w in ln) for ln in lines)
    shown = ", ".join(f"{_amount(a)} {c}" for a, c in got) or "ничего"
    return f"нет: прочитал «{read}» → {shown}"


def main():
    import openpyxl
    path = ARGS[0]
    label = ARGS[1] if len(ARGS) > 1 else ""
    png = shoot()
    jobs = [(png, i, mode) for mode in ("eng", "auto") for i in range(len(CASES))]
    with ProcessPoolExecutor(max_workers=6, initializer=_init) as ex:
        res = list(ex.map(work, jobs, chunksize=4))
    eng, auto = res[:len(CASES)], res[len(CASES):]
    if path == "-":                 # только напечатать: сравнить версии кода (REPO=...)
        for n, (e, a) in enumerate(zip(eng, auto), 1):
            print(f"{n}\t{e}\t{a}")
        return

    wb = openpyxl.load_workbook(path)
    ws = wb["Проверка валют"]
    done = 0
    for row in range(1, ws.max_row + 1):
        n = ws.cell(row, 1).value
        if not isinstance(n, int) or not 1 <= n <= len(CASES):
            continue
        if ws.cell(row, 2).value != CASES[n - 1][0]:
            sys.exit(f"строка {row}: на листе «{ws.cell(row, 2).value}», "
                     f"в списке «{CASES[n - 1][0]}» — список и лист разошлись")
        ws.cell(row, 9).value = eng[n - 1]
        ws.cell(row, 10).value = auto[n - 1]
        done += 1
    for row in range(1, 12):                       # заголовки столбцов — с датой прогона
        if label and str(ws.cell(row, 9).value or "").startswith("Мой прогон"):
            ws.cell(row, 9).value = f"Мой прогон: eng (как у вас сейчас), {label}"
            ws.cell(row, 10).value = f"Мой прогон: «авто» (как у новых), {label}"
    name, ext = os.path.splitext(os.path.basename(path))
    backup = os.path.join(tempfile.gettempdir(), f"{name} (до прогона){ext}")
    shutil.copy2(path, backup)
    wb.save(path)
    bad_eng = sum(v != "ок" for v in eng)
    bad_auto = sum(v != "ок" for v in auto)
    print(f"записано строк: {done}; не прошло: eng {bad_eng}, авто {bad_auto} из {len(CASES)}")
    print("копия книги до записи:", backup)


if __name__ == "__main__":
    main()
